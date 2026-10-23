"""Генераторы тестовых документов (в тестах нет заранее подготовленных файлов)."""
import shutil
import subprocess
from pathlib import Path

import pymupdf

TEXT_LINES = ["Протокол № ТЕСТ-001", "Статья 12.9, превышение скорости"]


def make_text_pdf(path: Path, pages=(TEXT_LINES,)) -> Path:
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page()
        html = "".join(f"<p>{l}</p>" for l in lines)
        page.insert_htmlbox(pymupdf.Rect(50, 50, 540, 400), html)
    doc.save(str(path))
    doc.close()
    return path


def make_scan_pdf(path: Path, lines=TEXT_LINES, mixed_with_text=None) -> Path:
    """PDF, где страница — только картинка (как скан). mixed_with_text — доп. текстовая страница."""
    tmp = path.with_suffix(".src.pdf")
    make_text_pdf(tmp, (lines,))
    src = pymupdf.open(str(tmp))
    png = src[0].get_pixmap(dpi=300).tobytes("png")
    src.close()
    tmp.unlink()
    doc = pymupdf.open()
    if mixed_with_text:
        page = doc.new_page()
        page.insert_htmlbox(pymupdf.Rect(50, 50, 540, 400),
                            "".join(f"<p>{l}</p>" for l in mixed_with_text))
    page = doc.new_page()
    page.insert_image(page.rect, stream=png)
    doc.save(str(path))
    doc.close()
    return path


def make_blank_pdf(path: Path) -> Path:
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(path))
    return path


def make_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_paragraph("Заголовок документа")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Номер"
    t.cell(0, 1).text = "ТЕСТ-001"
    t.cell(1, 0).merge(t.cell(1, 1)).text = "Объединённая ячейка"
    d.add_paragraph("Последний абзац")
    d.save(str(path))
    return path


def make_empty_docx(path: Path) -> Path:
    import docx
    docx.Document().save(str(path))
    return path


def make_odt(path: Path) -> Path:
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableRow
    from odf.text import P
    doc = OpenDocumentText()
    doc.text.addElement(P(text="Абзац ODT"))
    table = Table()
    row = TableRow()
    for value in ("Номер", "ТЕСТ-002"):
        cell = TableCell()
        cell.addElement(P(text=value))
        row.addElement(cell)
    table.addElement(row)
    doc.text.addElement(table)
    doc.save(str(path))
    return path


def have_soffice() -> bool:
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def make_doc(path: Path, docx_path: Path) -> Path:
    """Создаёт .doc конвертацией DOCX через LibreOffice."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    out = path.parent
    subprocess.run([exe, "--headless", "--convert-to", "doc", "--outdir", str(out), str(docx_path)],
                   capture_output=True, timeout=240, check=True)
    produced = out / (docx_path.stem + ".doc")
    if produced != path:
        produced.replace(path)
    return path


WRAPPED_PARAGRAPHS = [
    "ПРОТОКОЛ 86 НН 000001", "об административном правонарушении", "в отношении юридического лица",
    "06 февраля 2026 г. 15 ч. 30 мин. г. Тестоград, ул. Тестовая, 1",
    "Начальник отделения дорожного надзора отдела ГИБДД УМВД России по г. Тестограду, майор полиции Тестов Т. Т.",
    "в присутствии защитника юридического лица:",
    "составил настоящий протокол за нарушение, предусмотренное ч. 1 ст. 12.34 КоАП РФ в отношении:",
    'МБУ "Тестовое управление"', "расположенного по адресу: г. Тестоград, ул. Тестовая д. 2",
    "ОБСТОЯТЕЛЬСТВА ПРАВОНАРУШЕНИЯ:",
    "По результатам контрольных (надзорных) действий – осмотра, инструментального обследования, "
    "проведённых с 11 ч. 34 мин. по 16 ч. 10 мин. 05.02.2026, в рамках проведения постоянного рейда, "
    "а именно нарушены требования ст. 12 Федерального закона от 10.12.1995 №196-ФЗ, п. 13 ОПД ПДД РФ, "
    "п. 8.1, 8.8 ГОСТ Р 50597-2017, выразившиеся:",
    "в формировании снежных валов непосредственно на границе пешеходных переходов по ул. Мира д. 95;",
    "в наличии ям на покрытии проезжей части по ул. Образцовая д. 2 в нарушение п. 5.2.4 ГОСТ Р 50597-2017.",
    "По итогам контрольных (надзорных) действий составлен 1 протокол осмотра.",
    "В соответствии с п. 3.18 ГОСТ Р 59434-2021 снежный вал – это слой снега.",
]


def make_wrapped_pdf(path: Path, width: float = 230) -> Path:
    """Текстовый PDF с узкой колонкой: строки переносятся внутри фраз, как в реальных PDF."""
    doc = pymupdf.open()
    page = doc.new_page()
    rc = page.insert_htmlbox(pymupdf.Rect(40, 40, 40 + width, 800),
                             "".join(f"<p>{x}</p>" for x in WRAPPED_PARAGRAPHS))
    assert rc[0] >= 0, "текст не поместился на страницу"
    doc.save(str(path))
    doc.close()
    return path
